from __future__ import annotations

from io import BytesIO
from typing import Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from models import FinalMatchResult, MatchPairScoreDetail, UnmatchedStudent, UnusedVolunteer


HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAD3")
SUB_HEADER_FILL = PatternFill(fill_type="solid", fgColor="EAD1DC")


def _style_header_row(ws, row_idx: int = 1):
    for cell in ws[row_idx]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_all_cells(ws):
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _set_column_widths(ws, min_width: int = 10, max_width: int = 48):
    # 根据内容粗略自动列宽
    for col_cells in ws.columns:
        max_len = 0
        col_idx = col_cells[0].column
        for cell in col_cells:
            try:
                v = "" if cell.value is None else str(cell.value)
            except Exception:
                v = ""
            # 中文粗略按字符数估计
            l = min(len(v), 200)
            if l > max_len:
                max_len = l
        width = max(min_width, min(max_width, int(max_len * 1.2) + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _prepare_ws(ws, title: str):
    ws.title = title
    ws.freeze_panes = "A2"


def _append_rows(ws, headers: List[str], rows: Iterable[List[object]]):
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_header_row(ws, 1)
    _style_all_cells(ws)
    _set_column_widths(ws)


def _round2(v):
    try:
        return round(float(v), 2)
    except Exception:
        return v


def build_matching_export_workbook_bytes(db: Session) -> bytes:
    """
    导出4个sheet：
    1. 最终匹配结果
    2. 未匹配学生
    3. 未使用老师
    4. 全部合法pair打分明细
    """
    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    # -----------------------
    # Sheet 1: 最终匹配结果
    # -----------------------
    ws1 = wb.create_sheet("最终匹配结果")
    _prepare_ws(ws1, "最终匹配结果")

    final_rows = db.scalars(
        select(FinalMatchResult).order_by(FinalMatchResult.total_score.desc(), FinalMatchResult.id.asc())
    ).all()

    final_headers = [
        "阶段",
        "学生序号", "学生姓名",
        "老师序号", "老师姓名",
        "辅导方式", "老师匹配模式",
        "总分",
        "年级分", "科目分", "主观匹配分(DS)", "特殊需求惩罚分(DS)",
        "DS主观评分理由摘要",
        "DS特殊需求惩罚理由摘要",
    ]
    final_data = []
    for r in final_rows:
        final_data.append([
            r.phase,
            r.student_seq, r.student_name,
            r.teacher_seq, r.teacher_name,
            r.tutoring_mode, r.teacher_match_mode,
            _round2(r.total_score),
            _round2(r.grade_score), _round2(r.subject_score),
            _round2(r.ds_subjective_score), _round2(r.ds_special_penalty),
            r.ds_subjective_reason_summary or "",
            r.ds_special_reason_summary or "",
        ])
    _append_rows(ws1, final_headers, final_data)

    # -----------------------
    # Sheet 2: 未匹配学生
    # -----------------------
    ws2 = wb.create_sheet("未匹配学生")
    _prepare_ws(ws2, "未匹配学生")

    unmatched_rows = db.scalars(
        select(UnmatchedStudent).order_by(UnmatchedStudent.id.asc())
    ).all()

    unmatched_headers = [
        "学生序号", "学生姓名", "性别", "学段", "线上/线下", "是否优先", "原因摘要"
    ]
    unmatched_data = []
    for r in unmatched_rows:
        unmatched_data.append([
            r.student_seq, r.student_name, r.gender, r.stage_code,
            r.tutoring_mode, 1 if r.is_priority else 0,
            r.reason_summary or "",
        ])
    _append_rows(ws2, unmatched_headers, unmatched_data)

    # -----------------------
    # Sheet 3: 未使用老师
    # -----------------------
    ws3 = wb.create_sheet("未使用老师")
    _prepare_ws(ws3, "未使用老师")

    unused_rows = db.scalars(
        select(UnusedVolunteer).order_by(UnusedVolunteer.id.asc())
    ).all()

    unused_headers = [
        "老师序号", "老师姓名", "性别", "辅导方式", "匹配模式", "原因摘要"
    ]
    unused_data = []
    for r in unused_rows:
        unused_data.append([
            r.teacher_seq, r.teacher_name, r.gender, r.tutoring_mode, r.match_mode,
            r.reason_summary or "",
        ])
    _append_rows(ws3, unused_headers, unused_data)

    # -----------------------
    # Sheet 4: 全部合法pair打分明细
    # -----------------------
    ws4 = wb.create_sheet("合法Pair打分明细")
    _prepare_ws(ws4, "合法Pair打分明细")

    pair_rows = db.scalars(
        select(MatchPairScoreDetail).order_by(MatchPairScoreDetail.total_score.desc(), MatchPairScoreDetail.id.asc())
    ).all()

    pair_headers = [
        "阶段",
        "学生序号", "学生姓名", "学生是否优先", "学生辅导方式",
        "老师序号", "老师姓名", "老师匹配模式", "老师辅导方式",
        "总分",
        "年级分", "科目分", "科目原始最佳分(0-9)",
        "主观匹配分(DS)", "主观评分状态", "主观缓存命中", "主观理由摘要",
        "特殊需求惩罚分(DS)", "特殊评分状态", "特殊缓存命中", "特殊理由摘要",
        "是否入选最终结果",
    ]
    pair_data = []
    for r in pair_rows:
        pair_data.append([
            r.phase,
            r.student_seq, r.student_name, 1 if r.student_is_priority else 0, r.student_tutoring_mode,
            r.teacher_seq, r.teacher_name, r.teacher_match_mode, r.teacher_tutoring_mode,
            _round2(r.total_score),
            _round2(r.grade_score), _round2(r.subject_score), _round2(r.best_subject_raw),
            _round2(r.ds_subjective_score), r.ds_subjective_status, 1 if r.ds_subjective_cache_hit else 0, r.ds_subjective_reason_summary or "",
            _round2(r.ds_special_penalty), r.ds_special_status, 1 if r.ds_special_cache_hit else 0, r.ds_special_reason_summary or "",
            1 if r.selected_in_final else 0,
        ])
    _append_rows(ws4, pair_headers, pair_data)

    # 额外信息说明（可选）
    if ws4.max_row >= 1:
        ws4["A1"].comment = None  # 保持简洁

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()